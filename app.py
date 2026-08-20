import os
import uuid
import shutil
import logging
import threading
from pathlib import Path
from typing import Optional
from fastapi import FastAPI, Form, UploadFile, File, HTTPException, BackgroundTasks
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

# Tự động đọc file .env ở local nếu có
env_path = Path(".env")
if env_path.exists():
    with open(env_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ[k.strip()] = v.strip()

# Import transcriber
from transcriber import VideoTranscriber

# Logging
logging.basicConfig(level=logging.INFO)
log = logging.getLogger("web-service")

app = FastAPI(title="Video-to-Text Transcription Service")

# Whisper model cache
whisper_models_cache = {}
whisper_cache_lock = threading.Lock()

def get_cached_whisper_model(model_size: str):
    """
    Truy xuất mô hình Whisper từ cache hoặc tải mới nếu chưa có.
    Đảm bảo an toàn đa luồng bằng Lock.
    """
    with whisper_cache_lock:
        if model_size not in whisper_models_cache:
            log.info(f"Mô hình Whisper size '{model_size}' chưa có trong cache. Tiến hành tải...")
            try:
                from faster_whisper import WhisperModel
            except ImportError:
                raise ImportError("Vui lòng cài đặt: pip install faster-whisper")
            
            # Khởi tạo Whisper Model với num_workers=1 để CPU tự tối ưu tính toán
            model = WhisperModel(
                model_size,
                device="cpu",
                compute_type="int8",
                num_workers=1
            )
            whisper_models_cache[model_size] = model
            log.info(f"Đã lưu mô hình Whisper size '{model_size}' vào cache thành công.")
        else:
            log.info(f"Tìm thấy mô hình Whisper size '{model_size}' trong cache. Sử dụng lại (0s tải từ đĩa).")
        return whisper_models_cache[model_size]

# Global state to keep track of tasks
# task_id -> { "status": "processing/completed/failed", "filename": str, "progress": int, "total": int, "text": str, "error": str }
tasks_db = {}

# Lưu trữ trạng thái sử dụng của Groq API Key
groq_keys_status = {}

def init_groq_keys():
    api_keys_str = os.getenv("GROQ_API_KEY", "") or os.getenv("ROG_API_KEY", "")
    # Thay thế xuống dòng bằng dấu phẩy và phân tách các keys
    keys = [k.strip() for k in api_keys_str.replace("\n", ",").replace("\r", ",").split(",") if k.strip()]
    # Dọn dẹp các key không còn cấu hình
    for k in list(groq_keys_status.keys()):
        if k not in keys:
            del groq_keys_status[k]
    # Khởi tạo key mới
    for k in keys:
        if k not in groq_keys_status:
            masked = f"{k[:7]}...{k[-4:]}" if len(k) > 10 else "Invalid Key"
            groq_keys_status[k] = {
                "masked": masked,
                "used_seconds": 0.0,
                "status": "Active",
                "limit_seconds": 3600.0 # 60 phút
            }

def update_key_status(key, status, used_seconds=0.0):
    init_groq_keys()
    if key in groq_keys_status:
        groq_keys_status[key]["status"] = status
        groq_keys_status[key]["used_seconds"] += used_seconds
        if groq_keys_status[key]["used_seconds"] >= groq_keys_status[key]["limit_seconds"]:
            groq_keys_status[key]["status"] = "Rate Limited"
        log.info(f"Cập nhật key {groq_keys_status[key]['masked']}: status={status}, used_seconds={used_seconds}")

# Ensure required directories exist
UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

# Path to the model
MODEL_PATH = Path("models") / "vosk-model-small-vn-0.4"


def run_transcription_task(task_id: str, video_path: str, filename: str, engine: str, workers: int, is_temp_file: bool, whisper_model: str = "base", language: str = "vi"):
    """
    Chạy tác vụ trích xuất và nhận dạng giọng nói trong luồng nền.
    """
    try:
        tasks_db[task_id]["status"] = "processing"
        
        # Callback cập nhật tiến trình
        def progress_callback(done, total):
            tasks_db[task_id]["progress"] = done
            tasks_db[task_id]["total"] = total

        # Cấu hình model Vosk tự động tìm thư mục phù hợp theo ngôn ngữ vi hoặc en
        model_dir = None
        if engine == "vosk":
            models_path = Path("models")
            if models_path.exists():
                for path in models_path.iterdir():
                    if path.is_dir():
                        name_lower = path.name.lower()
                        if language == "vi" and ("vn" in name_lower or "vietnamese" in name_lower):
                            model_dir = str(path)
                            break
                        elif language == "en" and ("en" in name_lower or "english" in name_lower):
                            model_dir = str(path)
                            break
            # Fallback nếu không quét được thư mục tự động
            if not model_dir:
                fallback_path = models_path / ("vosk-model-small-vn-0.4" if language == "vi" else "vosk-model-small-en-us-0.15")
                if fallback_path.exists():
                    model_dir = str(fallback_path)
        
        whisper_model_instance = None
        if engine == "whisper":
            whisper_model_instance = get_cached_whisper_model(whisper_model)
        elif engine == "groq":
            # Đọc danh sách API Keys ngăn cách bằng dấu phẩy hoặc xuống dòng
            api_keys_str = os.getenv("GROQ_API_KEY", "") or os.getenv("ROG_API_KEY", "")
            whisper_model_instance = [k.strip() for k in api_keys_str.replace("\n", ",").replace("\r", ",").split(",") if k.strip()]
            
        transcriber = VideoTranscriber(
            model_path=model_dir,
            language=language,
            max_workers=workers,
            chunk_duration=60,
            whisper_model_size=whisper_model,
            whisper_model_instance=whisper_model_instance
        )
        if engine == "groq":
            transcriber.key_status_callback = update_key_status

        log.info(f"Bắt đầu dịch task {task_id}: {filename}")
        
        audio_filename = f"{task_id}_audio.wav"
        audio_path = UPLOAD_DIR / audio_filename

        result = transcriber.transcribe(
            video_path,
            progress_callback=progress_callback,
            engine=engine,
            save_audio_path=str(audio_path),
            return_segments=True
        )
        
        # Hoàn thành
        tasks_db[task_id]["status"] = "completed"
        tasks_db[task_id]["text"] = result["text"]
        tasks_db[task_id]["segments"] = result["segments"]
        tasks_db[task_id]["audio_url"] = f"/api/audio/{task_id}"
        tasks_db[task_id]["progress"] = tasks_db[task_id]["total"] # Đảm bảo 100%
        log.info(f"Hoàn thành task {task_id} thành công ✓")

    except Exception as e:
        log.error(f"Lỗi khi dịch task {task_id}: {e}")
        tasks_db[task_id]["status"] = "failed"
        tasks_db[task_id]["error"] = str(e)
    
    finally:
        # Xóa file video tạm nếu được upload lên
        if is_temp_file and os.path.exists(video_path):
            try:
                os.remove(video_path)
                log.info(f"Đã xóa file tạm: {video_path}")
            except Exception as ex:
                log.warning(f"Không thể xóa file tạm {video_path}: {ex}")


@app.get("/", response_class=HTMLResponse)
async def get_index():
    """Trả về giao diện HTML chính."""
    index_path = Path("templates") / "index.html"
    if not index_path.exists():
        raise HTTPException(status_code=404, detail="Không tìm thấy index.html")
    
    with open(index_path, "r", encoding="utf-8") as f:
        return HTMLResponse(content=f.read())


@app.post("/api/transcribe")
async def start_transcription(
    background_tasks: BackgroundTasks,
    video_path: Optional[str] = Form(None),
    engine: str = Form("vosk"),
    workers: int = Form(4),
    whisper_model: str = Form("base"),
    language: str = Form("vi"),
    file: Optional[UploadFile] = File(None)
):
    """Bắt đầu tác vụ chuyển đổi video."""
    task_id = str(uuid.uuid4())
    is_temp_file = False
    target_path = ""
    filename = ""

    if video_path:
        # Sử dụng đường dẫn cục bộ trực tiếp
        resolved_path = Path(video_path).resolve()
        if not resolved_path.exists():
            return JSONResponse(status_code=400, content={"detail": f"Đường dẫn file không hợp lệ hoặc không tồn tại: {video_path}"})
        target_path = str(resolved_path)
        filename = resolved_path.name
    elif file:
        # Xử lý file upload lên server
        filename = file.filename
        target_path = str(UPLOAD_DIR / f"{task_id}_{filename}")
        with open(target_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        is_temp_file = True
    else:
        return JSONResponse(status_code=400, content={"detail": "Vui lòng nhập đường dẫn file cục bộ hoặc tải lên một file video!"})

    # Tạo record task trong DB tạm thời
    tasks_db[task_id] = {
        "status": "pending",
        "filename": filename,
        "progress": 0,
        "total": 0,
        "text": "",
        "segments": [],
        "audio_url": "",
        "error": ""
    }

    # Đưa tác vụ vào chạy nền
    background_tasks.add_task(
        run_transcription_task,
        task_id=task_id,
        video_path=target_path,
        filename=filename,
        engine=engine,
        workers=workers,
        is_temp_file=is_temp_file,
        whisper_model=whisper_model,
        language=language
    )

    return {"task_id": task_id, "status": "pending"}


@app.get("/api/status/{task_id}")
async def get_task_status(task_id: str):
    """Lấy trạng thái tiến trình của tác vụ."""
    if task_id not in tasks_db:
        raise HTTPException(status_code=404, detail="Không tìm thấy task")
    return tasks_db[task_id]


@app.get("/api/quota")
async def get_quota_status():
    """API lấy hạn mức sử dụng hiện tại của các Groq API Keys."""
    init_groq_keys()
    keys_info = []
    for k, status in groq_keys_status.items():
        remaining = max(0.0, status["limit_seconds"] - status["used_seconds"])
        keys_info.append({
            "masked": status["masked"],
            "remaining_seconds": round(remaining, 2),
            "limit_seconds": status["limit_seconds"],
            "status": status["status"]
        })
    return {"keys": keys_info}


@app.get("/api/audio/{task_id}")
async def get_audio(task_id: str):
    """Phục vụ file âm thanh WAV đã trích xuất của task."""
    audio_path = UPLOAD_DIR / f"{task_id}_audio.wav"
    if not audio_path.exists():
        raise HTTPException(status_code=404, detail="Không tìm thấy tệp âm thanh")
    return FileResponse(audio_path, media_type="audio/wav")


@app.get("/api/export/docx/{task_id}")
async def export_docx(task_id: str):
    """Xuất văn bản nhận dạng dưới dạng file Word (.docx)."""
    if task_id not in tasks_db:
        raise HTTPException(status_code=404, detail="Không tìm thấy thông tin task")
        
    task = tasks_db[task_id]
    if task["status"] != "completed":
        raise HTTPException(status_code=400, detail="Task chưa hoàn thành để xuất file")
        
    filename = task["filename"]
    segments = task.get("segments", [])
    
    import docx
    import time
    from io import BytesIO
    
    doc = docx.Document()
    doc.add_heading(f"TRANSCRIPT — {filename}", 0)
    doc.add_paragraph(f"Tệp gốc: {filename}")
    doc.add_paragraph(f"Thời gian tạo: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("-" * 40)
    
    if segments:
        for seg in segments:
            # Tính toán timestamp format [mm:ss]
            start_sec = seg.get("start", 0)
            mins = int(start_sec // 60)
            secs = int(start_sec % 60)
            time_str = f"[{mins:02d}:{secs:02d}]"
            
            # Lấy text tương ứng (ưu tiên ghép từ words nếu có để đồng bộ 100% với UI)
            if "words" in seg and seg["words"]:
                seg_text = " ".join(w["word"] for w in seg["words"])
            else:
                seg_text = seg.get("text", "")
                
            p = doc.add_paragraph()
            # Thêm timestamp dạng in đậm
            run_time = p.add_run(f"{time_str}  ")
            run_time.bold = True
            
            # Thêm văn bản
            p.add_run(seg_text)
    else:
        # Fallback nếu không có dữ liệu segments
        text = task.get("text", "")
        for line in text.split("\n"):
            line = line.strip()
            if line:
                doc.add_paragraph(line)
                
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    # Chuẩn hóa tên file tải về
    safe_stem = Path(filename).stem
    # Loại bỏ ký tự không hợp lệ trong header content-disposition
    safe_stem = safe_stem.encode("ascii", "ignore").decode("ascii").replace(" ", "_")
    if not safe_stem:
        safe_stem = "transcript"
    docx_name = f"{safe_stem}_transcript.docx"
    
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={docx_name}"}
    )


if __name__ == "__main__":
    import uvicorn
    import threading
    import webbrowser

    print("\n" + "="*80)
    print("  DỊCH VỤ CHUYỂN VIDEO THÀNH TEXT OFFLINE ĐANG KHỞI CHẠY...")
    print("  Trình duyệt web sẽ tự động mở sau vài giây...")
    print("  Nếu không tự mở, hãy truy cập: http://127.0.0.1:8000")
    print("="*80 + "\n")

    # Tự động mở trình duyệt mặc định sau 1.5 giây
    threading.Timer(1.5, lambda: webbrowser.open("http://127.0.0.1:8000")).start()

    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=False)

